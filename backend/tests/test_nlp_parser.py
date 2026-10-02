import io
import pytest
import docx

from app.models.candidate import DegreeLevel
from app.services.nlp.extractor import clean_text, extract_text_from_docx
from app.services.nlp.parser import (
    calculate_tenure_and_roles,
    chunk_sections,
    detect_highest_degree,
    extract_contact_info,
    parse_resume,
    scrub_pii_for_blind_screening,
)
from app.services.nlp.taxonomy import skills_extractor


def test_clean_text_normalizes_whitespace_and_bullets():
    dirty = "Hello \r\n\r\n World\t• Bullet 1\u2022 Bullet 2   Extra   Spaces\x00\x08"
    cleaned = clean_text(dirty)
    assert "\r" not in cleaned
    assert "\x00" not in cleaned
    assert "• Bullet 1" in cleaned
    assert "Bullet 2" in cleaned
    assert "Extra Spaces" in cleaned


def test_docx_text_extraction():
    doc = docx.Document()
    doc.add_paragraph("Alice Smith")
    doc.add_paragraph("Senior Python Developer")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills:"
    table.rows[0].cells[1].text = "FastAPI, Docker, Kubernetes"

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    extracted = extract_text_from_docx(file_bytes)
    assert "Alice Smith" in extracted
    assert "Senior Python Developer" in extracted
    assert "FastAPI, Docker, Kubernetes" in extracted


def test_skills_taxonomy_normalization():
    text = "Proficient in K8s, react.js, PostgreSQL, Golang, AWS Lambda, and PyTorch."
    extracted = skills_extractor.extract_skills(text)
    assert "Kubernetes" in extracted
    assert "React" in extracted
    assert "PostgreSQL" in extracted
    assert "Go" in extracted
    assert "AWS" in extracted
    assert "PyTorch" in extracted


def test_contact_info_extraction():
    text = """
    Bob Vance
    Scranton, PA
    Email: bob.vance@vancerefrigeration.com
    Phone: (555) 867-5309
    LinkedIn: https://www.linkedin.com/in/bobvance
    GitHub: https://github.com/bobvance
    """
    contact = extract_contact_info(text)
    assert contact["email"] == "bob.vance@vancerefrigeration.com"
    assert contact["phone"] == "(555) 867-5309"
    assert "linkedin.com/in/bobvance" in contact["linkedin_url"]
    assert "github.com/bobvance" in contact["portfolio_url"]
    assert contact["full_name"] == "Bob Vance"


def test_tenure_and_role_calculation():
    experience_text = """
    Lead Architect - Wayne Enterprises
    Jan 2021 - Present
    Designed cloud microservices.

    Software Engineer | LexCorp
    Jan 2018 - Dec 2020
    Maintained backend services.
    """
    total_months, title, company, roles = calculate_tenure_and_roles(experience_text)
    # 2018 to 2020 is 36 months + 2021 to 2026 is ~60+ months => total > 90 months
    assert total_months >= 90
    assert len(roles) == 2
    assert "Architect" in str(title)


def test_highest_degree_detection():
    phd_text = "Doctor of Philosophy (Ph.D.) in Computer Science, MIT 2019"
    deg, records = detect_highest_degree(phd_text)
    assert deg == DegreeLevel.PHD

    masters_text = "Master of Science in Software Engineering, Carnegie Mellon 2021"
    deg_m, _ = detect_highest_degree(masters_text)
    assert deg_m == DegreeLevel.MASTERS

    bachelors_text = "Bachelor of Science in Information Technology, 2015"
    deg_b, _ = detect_highest_degree(bachelors_text)
    assert deg_b == DegreeLevel.BACHELORS


def test_blind_screening_pii_scrubber():
    profile = {
        "full_name": "Tony Stark",
        "email": "tony@starkindustries.com",
        "phone": "+1-555-0199",
        "location": "Malibu, CA",
        "skills": ["Python", "C++", "Robotics"],
        "education": [{"degree": "BACHELORS", "institution": "MIT", "graduation_year": 1991}],
    }
    blind = scrub_pii_for_blind_screening(profile)
    assert blind["full_name"] == "[REDACTED FOR BLIND SCREENING]"
    assert blind["email"] == "[REDACTED]"
    assert blind["phone"] == "[REDACTED]"
    assert "graduation_year" not in blind["education"][0]
    assert blind["skills"] == ["Python", "C++", "Robotics"]


def test_full_parse_resume():
    resume = """
    Bruce Banner
    Email: bruce.banner@avengers.org
    Phone: (555) 999-0000
    LinkedIn: linkedin.com/in/brucebanner

    PROFESSIONAL EXPERIENCE
    Senior Research Scientist - Culver University
    June 2019 - Present
    - Developed advanced simulation models using Python and PyTorch.
    - Managed Docker containers and deployed on Kubernetes (K8s).

    EDUCATION
    Ph.D. in Nuclear Physics - Harvard University
    2015 - 2019

    TECHNICAL SKILLS
    Python, C++, Docker, Kubernetes, PyTorch, Linux, Git
    """
    parsed = parse_resume(resume)
    assert parsed["contact"]["email"] == "bruce.banner@avengers.org"
    assert parsed["highest_degree"] == "PHD"
    assert "Python" in parsed["skills"]
    assert "Kubernetes" in parsed["skills"]
    assert "PyTorch" in parsed["skills"]
    assert parsed["total_tenure_months"] > 60
    assert parsed["blind_view"]["full_name"] == "[REDACTED FOR BLIND SCREENING]"
