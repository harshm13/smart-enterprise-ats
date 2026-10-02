import io
import uuid
import pytest
import docx
from httpx import AsyncClient
from jose import jwt
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.user import User, UserRole


def _make_jwt(tenant_id: str, sub: str, role: str = "recruiter") -> str:
    return jwt.encode(
        {"sub": sub, "org_id": tenant_id, "role": role},
        key="secret",
        algorithm="HS256",
    )


async def _create_tenant_and_user(
    client: AsyncClient,
    db_session: AsyncSession,
    slug: str,
    name: str,
    role: UserRole = UserRole.RECRUITER,
) -> tuple[dict, User, str]:
    sub = f"auth0|{uuid.uuid4().hex[:8]}"
    r = await client.post("/v1/tenants/", json={"slug": slug, "name": name})
    assert r.status_code == 201
    tenant_data = r.json()
    tenant_id = uuid.UUID(tenant_data["id"])

    user = User(
        tenant_id=tenant_id,
        email=f"recruiter@{slug}.com",
        full_name=f"{name} Recruiter",
        role=role,
        auth0_sub=sub,
    )
    db_session.add(user)
    await db_session.flush()
    await db_session.refresh(user)

    token = _make_jwt(str(tenant_id), sub, role.value)
    return tenant_data, user, token


def _build_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Diana Prince")
    doc.add_paragraph("Email: diana.prince@themyscira.gov")
    doc.add_paragraph("Phone: (555) 777-8888")
    doc.add_paragraph("LinkedIn: linkedin.com/in/dianaprince")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Principal Distributed Systems Engineer | Amazon Web Services")
    doc.add_paragraph("Jan 2020 - Present")
    doc.add_paragraph("Engineered high-scale microservices using Go, Kubernetes, and Docker.")
    doc.add_paragraph("Senior Backend Engineer | Oracle")
    doc.add_paragraph("Jan 2016 - Dec 2019")
    doc.add_paragraph("Built database engines with Python and PostgreSQL.")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("Master of Science in Computer Science | MIT 2015")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Go, Python, Kubernetes, Docker, PostgreSQL, AWS, Microservices, Redis")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_upload_resume_unauthenticated_fails(client: AsyncClient, db_session: AsyncSession) -> None:
    """Uploading resume without auth header returns 401."""
    files = {"file": ("resume.docx", b"dummy content", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = await client.post("/v1/resumes/upload", files=files)
    assert res.status_code == 401


@pytest.mark.asyncio
async def test_upload_resume_invalid_extension(client: AsyncClient, db_session: AsyncSession) -> None:
    """Uploading unsupported file formats returns 415."""
    _, _, token = await _create_tenant_and_user(client, db_session, slug="ext-corp", name="Ext Corp")
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("resume.exe", b"malicious binary", "application/octet-stream")}
    res = await client.post("/v1/resumes/upload", files=files, headers=headers)
    assert res.status_code == 415


@pytest.mark.asyncio
async def test_upload_resume_and_parse_end_to_end(client: AsyncClient, db_session: AsyncSession) -> None:
    """Uploads a valid DOCX resume, verifies 202 Accepted, and verifies candidate parsing."""
    _, user, token = await _create_tenant_and_user(
        client, db_session, slug="apex-corp", name="Apex Corp"
    )
    headers = {"Authorization": f"Bearer {token}"}

    docx_bytes = _build_sample_docx()
    files = {
        "file": (
            "diana_prince_resume.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }

    # 1. Upload resume
    res = await client.post("/v1/resumes/upload", files=files, headers=headers)
    assert res.status_code == 202
    data = res.json()
    assert "task_id" in data
    assert "candidate_id" in data
    candidate_id = data["candidate_id"]

    # 2. Poll task status
    task_id = data["task_id"]
    task_res = await client.get(f"/v1/resumes/tasks/{task_id}", headers=headers)
    assert task_res.status_code == 200

    # 3. Retrieve Candidate Profile
    cand_res = await client.get(f"/v1/candidates/{candidate_id}", headers=headers)
    assert cand_res.status_code == 200
    candidate = cand_res.json()
    assert candidate["id"] == candidate_id
    assert candidate["email"] == "diana.prince@themyscira.gov"
    assert candidate["full_name"] == "Diana Prince"
    assert candidate["highest_degree"] == "MASTERS"
    assert "Kubernetes" in candidate["skills"]
    assert "Python" in candidate["skills"]
    assert "PostgreSQL" in candidate["skills"]
    assert candidate["total_tenure_months"] > 60

    # 4. Retrieve Blind Screening View (Anti-bias)
    blind_res = await client.get(f"/v1/candidates/{candidate_id}?blind=true", headers=headers)
    assert blind_res.status_code == 200
    blind_cand = blind_res.json()
    assert blind_cand["full_name"] == "[REDACTED FOR BLIND SCREENING]"
    assert blind_cand["email"] == "[REDACTED]"
    assert blind_cand["phone"] == "[REDACTED]"
    assert "Kubernetes" in blind_cand["skills"]

    # 5. Filter candidates by skill
    skill_filter_res = await client.get("/v1/candidates/?skills=Kubernetes,Go", headers=headers)
    assert skill_filter_res.status_code == 200
    assert skill_filter_res.json()["total"] >= 1

    # 6. Delete candidate (GDPR right to be forgotten)
    del_res = await client.delete(f"/v1/candidates/{candidate_id}", headers=headers)
    assert del_res.status_code == 204

    # Candidate should now be gone
    gone_res = await client.get(f"/v1/candidates/{candidate_id}", headers=headers)
    assert gone_res.status_code == 404


@pytest.mark.asyncio
async def test_presigned_url_generation(client: AsyncClient, db_session: AsyncSession) -> None:
    """Generates pre-signed upload URL for client uploads."""
    _, _, token = await _create_tenant_and_user(client, db_session, slug="presign-corp", name="Presign Corp")
    headers = {"Authorization": f"Bearer {token}"}

    res = await client.post(
        "/v1/resumes/presigned-url",
        json={"file_name": "candidate.pdf", "mime_type": "application/pdf"},
        headers=headers,
    )
    assert res.status_code == 200
    data = res.json()
    assert "upload_url" in data
    assert "key" in data
    assert data["headers"]["Content-Type"] == "application/pdf"


@pytest.mark.asyncio
async def test_candidate_cross_tenant_isolation(client: AsyncClient, db_session: AsyncSession) -> None:
    """Recruiter in Tenant A cannot view or delete Candidate in Tenant B."""
    _, _, token_a = await _create_tenant_and_user(client, db_session, slug="tenant-a-resumes", name="Tenant A")
    headers_a = {"Authorization": f"Bearer {token_a}"}

    _, _, token_b = await _create_tenant_and_user(client, db_session, slug="tenant-b-resumes", name="Tenant B")
    headers_b = {"Authorization": f"Bearer {token_b}"}

    # Upload candidate in Tenant A
    docx_bytes = _build_sample_docx()
    files = {"file": ("resume_a.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    upload_a = await client.post("/v1/resumes/upload", files=files, headers=headers_a)
    assert upload_a.status_code == 202
    cand_a_id = upload_a.json()["candidate_id"]

    # Recruiter B tries to GET Candidate A -> 404
    get_b = await client.get(f"/v1/candidates/{cand_a_id}", headers=headers_b)
    assert get_b.status_code == 404

    # Recruiter B tries to DELETE Candidate A -> 404
    del_b = await client.delete(f"/v1/candidates/{cand_a_id}", headers=headers_b)
    assert del_b.status_code == 404

    # Candidate A is not present in Tenant B candidate list
    list_b = await client.get("/v1/candidates/", headers=headers_b)
    assert list_b.status_code == 200
    assert list_b.json()["total"] == 0
